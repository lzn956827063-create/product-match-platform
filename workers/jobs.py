"""Database-owned job state, bounded chunks and fencing for at-least-once delivery."""
import csv
import io
import logging
import random
import threading
import time
from functools import lru_cache

from openpyxl import Workbook
from sqlalchemy import func, select

from packages.domain import storage
from packages.domain.db import now, transaction, uid
from packages.domain.models import Candidate, CatalogVersion, Chunk, Export, ExportRow, Item, Outbox, Product, Run, Scheduler, Source
from packages.matching.engine import FEATURE_VERSION, INDEX_VERSION, Matcher
from packages.matching.normalize import RULE_VERSION, digest

log = logging.getLogger("jobs")
LEASE_SECONDS = 120
TERMINAL = {"SUCCEEDED", "FAILED", "CANCELLED"}


def complete_outbox(s, kind, ident):
    event = s.scalar(select(Outbox).where(Outbox.event_key == f"{kind}:{ident}"))
    if event:
        event.completed = True


def cancel_locked(s, run):
    # Fencing every chunk makes all old workers unable to commit, even if alive.
    for c in s.scalars(select(Chunk).where(Chunk.org_id == run.org_id, Chunk.run_id == run.id).with_for_update()):
        c.fence_token += 1
        c.lease_until = 0
        if c.status != "DONE":
            c.status = "CANCELLED"
    run.status, run.finished_at = "CANCELLED", now()
    complete_outbox(s, "match", run.id)


def start_run(run_id):
    with transaction(write=True) as s:
        s.scalar(select(Scheduler).where(Scheduler.id == "global").with_for_update())
        run = s.scalar(select(Run).where(Run.id == run_id).with_for_update())
        if not run or run.status in TERMINAL:
            if run:
                complete_outbox(s, "match", run.id)
            return None
        if run.status == "CANCEL_REQUESTED":
            cancel_locked(s, run)
            return None
        if run.status == "QUEUED":
            active = s.scalars(select(Run).where(Run.status.in_(["RUNNING", "CANCEL_REQUESTED"]))).all()
            if len(active) >= 2 or any(r.org_id == run.org_id for r in active):
                return None
            run.status, run.started_at = "RUNNING", now()
        return {"id": run.id, "org_id": run.org_id, "catalog_version_id": run.catalog_version_id, "manifest": run.manifest}


def claim_chunk(run_id):
    with transaction(write=True) as s:
        run = s.scalar(select(Run).where(Run.id == run_id).with_for_update())
        if run.status == "CANCEL_REQUESTED":
            cancel_locked(s, run)
            return None
        if run.status != "RUNNING":
            return None
        chunks = s.scalars(select(Chunk).where(Chunk.org_id == run.org_id, Chunk.run_id == run_id).order_by(Chunk.number).with_for_update()).all()
        for c in chunks:
            if c.status == "DONE" or c.retry_after > time.time() or (c.status == "RUNNING" and c.lease_until > time.time()):
                continue
            c.status, c.lease_until, c.fence_token, c.attempts = "RUNNING", time.time()+LEASE_SECONDS, c.fence_token+1, c.attempts+1
            return {"id": c.id, "org_id": c.org_id, "run_id": run.id, "source_ids": c.source_ids, "token": c.fence_token, "attempts": c.attempts}
        if all(c.status == "DONE" for c in chunks):
            run.status, run.finished_at = "SUCCEEDED", now()
            complete_outbox(s, "match", run.id)
        return None


def renew(claim):
    with transaction(write=True) as s:
        run = s.scalar(select(Run).where(Run.id == claim["run_id"], Run.org_id == claim["org_id"]).with_for_update())
        c = s.scalar(select(Chunk).where(Chunk.id == claim["id"], Chunk.org_id == claim["org_id"]).with_for_update())
        if run.status != "RUNNING" or c.fence_token != claim["token"] or c.lease_until <= time.time() or c.status != "RUNNING":
            return False
        c.lease_until = time.time()+LEASE_SECONDS
        return True


def commit_chunk(claim, results, seconds):
    with transaction(write=True) as s:
        run = s.scalar(select(Run).where(Run.id == claim["run_id"], Run.org_id == claim["org_id"]).with_for_update())
        c = s.scalar(select(Chunk).where(Chunk.id == claim["id"], Chunk.org_id == claim["org_id"]).with_for_update())
        if run.status != "RUNNING" or c.fence_token != claim["token"] or c.status != "RUNNING" or c.lease_until <= time.time():
            return False
        if {r[0] for r in results} != set(c.source_ids):
            raise ValueError("CHUNK_RESULT_INCOMPLETE")
        for source_id, suggestion, candidates in results:
            item = Item(id=uid(), org_id=run.org_id, run_id=run.id, source_id=source_id, suggestion=suggestion)
            s.add(item)
            s.flush()
            for candidate in candidates:
                s.add(Candidate(org_id=run.org_id, item_id=item.id, **candidate))
        c.status, c.lease_until = "DONE", 0
        run.processed += len(results)
        run.timings = {**run.timings, "matching_seconds": run.timings.get("matching_seconds", 0) + seconds}
        s.flush()
        remaining = s.scalar(select(func.count()).select_from(Chunk).where(Chunk.org_id == run.org_id, Chunk.run_id == run.id, Chunk.status != "DONE"))
        if not remaining:
            run.status, run.finished_at = "SUCCEEDED", now()
            complete_outbox(s, "match", run.id)
        return True


def fail_chunk(claim, exc):
    with transaction(write=True) as s:
        run = s.scalar(select(Run).where(Run.id == claim["run_id"]).with_for_update())
        c = s.scalar(select(Chunk).where(Chunk.id == claim["id"]).with_for_update())
        if c.fence_token != claim["token"] or run.status != "RUNNING":
            return
        if isinstance(exc, (ValueError, TypeError)) or c.attempts >= 4:
            run.status, run.error, run.finished_at = "FAILED", type(exc).__name__, now()
            run.failed = run.total - run.processed
            c.status = "FAILED"
            complete_outbox(s, "match", run.id)
        else:
            c.status, c.lease_until = "PENDING", 0
            c.retry_after = time.time() + [10, 30, 90][c.attempts-1] + random.random()*3


@lru_cache(maxsize=2)
def get_matcher(org_id, version_id, content_hash, policy_json):
    import json
    with transaction() as s:
        products = s.scalars(select(Product).where(Product.org_id == org_id, Product.version_id == version_id).order_by(Product.id)).all()
        data = [{"id": p.id, "normalized": p.normalized} for p in products]
    return Matcher(data, json.loads(policy_json))


def process_run(run_id):
    import json
    definition = start_run(run_id)
    if not definition:
        return
    t0 = time.perf_counter()
    try:
        if definition["manifest"]["rule_version"] != RULE_VERSION or definition["manifest"]["index_version"] != INDEX_VERSION or definition["manifest"]["feature_schema"] != FEATURE_VERSION:
            raise ValueError("UNSUPPORTED_MANIFEST_VERSION")
        matcher = get_matcher(definition["org_id"], definition["catalog_version_id"], definition["manifest"]["catalog_hash"], json.dumps(definition["manifest"]["policy"], sort_keys=True))
    except Exception as exc:
        with transaction(write=True) as s:
            run = s.scalar(select(Run).where(Run.id == run_id).with_for_update())
            if run.status == "RUNNING":
                run.status, run.error, run.failed, run.finished_at = "FAILED", type(exc).__name__, run.total-run.processed, now()
                complete_outbox(s, "match", run_id)
        log.error("index_failed run_id=%s error_type=%s", run_id, type(exc).__name__)
        return
    with transaction(write=True) as s:
        run = s.scalar(select(Run).where(Run.id == run_id).with_for_update())
        run.timings = {**run.timings, "index_seconds": time.perf_counter()-t0, "index_hash": matcher.index_hash}
    while claim := claim_chunk(run_id):
        stopped = threading.Event()
        lost = threading.Event()

        def heartbeat():
            while not stopped.wait(20):
                try:
                    if not renew(claim):
                        lost.set()
                        return
                except Exception:
                    lost.set()
                    return

        thread = threading.Thread(target=heartbeat, daemon=True)
        thread.start()
        try:
            started = time.perf_counter()
            with transaction() as s:
                sources = s.scalars(select(Source).where(Source.org_id == claim["org_id"], Source.id.in_(claim["source_ids"]))).all()
                source_data = [(x.id, x.normalized) for x in sources]
            results = []
            for i, (source_id, normalized) in enumerate(source_data):
                if lost.is_set() or (i % 20 == 0 and not renew(claim)):
                    break
                tick = time.perf_counter()
                suggestion, candidates = matcher.match(normalized)
                if time.perf_counter()-tick > 30:
                    raise TimeoutError("INFERENCE_TIMEOUT")
                results.append((source_id, suggestion, candidates))
            if len(results) == len(source_data) and not lost.is_set():
                commit_chunk(claim, results, time.perf_counter()-started)
            else:
                break
        except Exception as exc:
            fail_chunk(claim, exc)
            log.error("chunk_failed run_id=%s chunk_id=%s error_type=%s", run_id, claim["id"], type(exc).__name__)
            break
        finally:
            stopped.set()
            thread.join(timeout=1)
    # A cancelled worker must close the state after relinquishing its claim.
    with transaction(write=True) as s:
        run = s.scalar(select(Run).where(Run.id == run_id).with_for_update())
        if run.status == "CANCEL_REQUESTED":
            cancel_locked(s, run)


def safe_text(value):
    text = str(value if value is not None else "")
    if text and (text[0] in "=+-@\t\r\n" or text.lstrip().startswith(("=", "+", "-", "@")) or ord(text[0]) < 32):
        return "'" + text
    return text


EXPORT_HEADERS = ["来源文件", "工作表", "原始行号", "来源编号", "编号类型", "原始名称", "标准编号", "标准名称", "差异说明", "审核状态", "审核原因", "审核人", "审核时间", "运行编号", "决定编号"]


def process_export(export_id):
    with transaction(write=True) as s:
        export = s.scalar(select(Export).where(Export.id == export_id).with_for_update())
        if not export or export.status in ("SUCCEEDED", "FAILED"):
            if export:
                complete_outbox(s, "export", export.id)
            return
        export.status = "RUNNING"
        rows = [r.data for r in s.scalars(select(ExportRow).where(ExportRow.org_id == export.org_id, ExportRow.export_id == export.id).order_by(ExportRow.row_no))]
        org_id, fmt, snapshot_hash = export.org_id, export.format, export.snapshot_hash
    try:
        if digest(rows) != snapshot_hash:
            raise ValueError("SNAPSHOT_HASH_MISMATCH")
        headers = list(dict.fromkeys(EXPORT_HEADERS + [k for r in rows for k in r]))
        if fmt == "csv":
            buff = io.StringIO(newline="")
            writer = csv.writer(buff)
            writer.writerow(headers)
            writer.writerows([[safe_text(r.get(h, "")) for h in headers] for r in rows])
            content = buff.getvalue().encode("utf-8-sig")
        else:
            from openpyxl.styles import Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = "匹配结果"
            ws.append(headers)
            for r in rows:
                ws.append([safe_text(r.get(h, "")) for h in headers])
            for row in ws:
                for cell in row:
                    cell.data_type = "s"
                    cell.number_format = "@"
            for c in ws[1]:
                c.font = Font(color="FFFFFF", bold=True)
                c.fill = PatternFill("solid", fgColor="3554D1")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 24
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            buff = io.BytesIO()
            wb.save(buff)
            content = buff.getvalue()
        key = storage.put(org_id, content, fmt)
        with transaction(write=True) as s:
            export = s.scalar(select(Export).where(Export.id == export_id).with_for_update())
            if export.status != "SUCCEEDED":
                export.object_key, export.file_hash, export.status = key, digest(content), "SUCCEEDED"
            complete_outbox(s, "export", export.id)
    except Exception as exc:
        with transaction(write=True) as s:
            export = s.scalar(select(Export).where(Export.id == export_id).with_for_update())
            if export.status != "SUCCEEDED":
                export.status, export.error = "FAILED", type(exc).__name__
            complete_outbox(s, "export", export.id)
        log.error("export_failed export_id=%s error_type=%s", export_id, type(exc).__name__)
