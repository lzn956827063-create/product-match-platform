import csv
import io
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from .errors import require, Problem
from ..matching.normalize import LABELS, normalize

MAX_BYTES = 20 * 1024 * 1024
FIELDS = {"sku": ["sku", "编号", "商品编号", "内部编号", "来源编号", "商品编码"], "name": ["name", "名称", "商品名称", "品名"], "brand": ["brand", "品牌"], "model": ["model", "型号"], "specs": ["specs", "规格", "规格描述"], "ram": ["ram", "运行内存", "内存"], "storage": ["storage", "存储容量", "容量"], "color": ["color", "颜色"], "region": ["region", "销售版本", "地区版本", "版本"], "pack_count": ["pack_count", "包装数量", "套装数量"], "price": ["price", "价格", "报价", "单价"], "currency": ["currency", "币种"]}


def parse_file(data, filename, encoding="utf-8"):
    require(len(data) <= MAX_BYTES, 413, "FILE_TOO_LARGE", "文件不能超过 20MB")
    suffix = Path(filename).suffix.lower()
    require(suffix in (".csv", ".xlsx"), 422, "INVALID_FILE", "仅支持 CSV 和 XLSX 文件")
    require(encoding in ("utf-8", "gb18030"), 422, "INVALID_ENCODING", "请选择 UTF-8 或 GB18030")
    try:
        if suffix == ".csv":
            require(not data.startswith(b"PK") and b"\x00" not in data, 422, "INVALID_FILE", "CSV 文件格式无效")
            content = data.decode("utf-8-sig" if encoding == "utf-8" else "gb18030")
            csv.field_size_limit(10001)
            rows = []
            for row_no, row in enumerate(csv.reader(io.StringIO(content)), 1):
                require(row_no <= 50100, 422, 'ROW_LIMIT', '文件行数过多')
                validate_matrix([row])
                rows.append(row)
            validate_matrix(rows)
            return {"CSV": {"rows": rows, "formulas": []}}
        require(data.startswith(b"PK\x03\x04"), 422, "INVALID_FILE", "XLSX 文件签名无效")
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            require(len(infos) <= 2000 and sum(i.file_size for i in infos) <= 100 * 1024 * 1024, 422, "XLSX_LIMIT", "工作簿解压体积超过 100MB 或结构过大")
            require("xl/workbook.xml" in z.namelist() and not any("vbaproject" in i.filename.lower() for i in infos), 422, "INVALID_FILE", "不支持含宏的工作簿")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        require(len(wb.sheetnames) <= 10, 422, "XLSX_LIMIT", "工作表不能超过 10 个")
        result = {}
        for ws in wb:
            require(ws.max_row is None or ws.max_row <= 50100, 422, "ROW_LIMIT", "工作表超过 50,100 行")
            require(ws.max_column is None or ws.max_column <= 100, 422, "COLUMN_LIMIT", "工作表不能超过 100 列")
            rows, formulas = [], []
            for row_no, row in enumerate(ws.iter_rows(), 1):
                vals = []
                for col_no, c in enumerate(row):
                    if c.data_type == "f":
                        formulas.append([row_no, col_no])
                    value = c.value
                    # Preserve formatted numeric identifiers such as 000123.
                    if isinstance(value, (int, float)) and not isinstance(value, bool) and value == int(value) and re.fullmatch(r"0+", c.number_format or ""):
                        value = str(int(value)).zfill(len(c.number_format))
                    vals.append("" if value is None else str(value))
                rows.append(vals)
                require(row_no <= 50100, 422, "ROW_LIMIT", "工作表行数过多")
            validate_matrix(rows)
            result[ws.title] = {"rows": rows, "formulas": formulas}
        wb.close()
        return result
    except Problem:
        raise
    except (UnicodeError, csv.Error, zipfile.BadZipFile, Exception) as exc:
        raise Problem(422, "PARSE_ERROR", "文件解析失败，请核对编码和文件格式") from exc


def validate_matrix(rows):
    require(len(rows) <= 50100, 422, "ROW_LIMIT", "文件行数过多")
    require(all(len(r) <= 100 for r in rows), 422, "COLUMN_LIMIT", "文件不能超过 100 列")
    require(all(len(c) <= 10000 for r in rows for c in r), 422, "CELL_LIMIT", "单元格不能超过 10,000 字符")


def preview(parsed, sheet, header_row):
    require(sheet in parsed, 422, "INVALID_SHEET", "工作表不存在")
    rows = parsed[sheet]["rows"]
    require(1 <= header_row <= min(len(rows), 100), 422, "INVALID_HEADER", "表头行必须在前 100 行内")
    headers = [str(x).strip() or f"未命名列{i+1}" for i, x in enumerate(rows[header_row - 1])]
    require(len(headers) == len(set(headers)), 422, "DUPLICATE_HEADERS", "表头列名重复，请先修正")
    mapping = {key: next((h for h in headers if h.lower() in aliases), None) for key, aliases in FIELDS.items()}
    return {"headers": headers, "rows": [{"row_no": i, "values": dict(zip(headers, row))} for i, row in enumerate(rows[header_row:], header_row + 1)][:50], "total": len(rows) - header_row, "suggested_mapping": {k: v for k, v in mapping.items() if v}, "formula_cells": parsed[sheet]["formulas"][:100]}


def _transform_fields(fields, transformations):
    require(set(transformations) <= set(FIELDS), 422, "INVALID_TRANSFORMATION", "转换规则包含未知标准字段")
    output = {k: str(v or "") for k, v in fields.items()}
    for field, rule in transformations.items():
        require(isinstance(rule, dict) and set(rule) <= {"trim", "case", "default", "replace"}, 422, "INVALID_TRANSFORMATION", "转换规则格式不正确")
        value = output.get(field, "")
        if rule.get("trim", True):
            value = value.strip()
        replacements = rule.get("replace", {})
        require(isinstance(replacements, dict) and len(replacements) <= 200, 422, "INVALID_TRANSFORMATION", "枚举替换规则过多或格式不正确")
        value = str(replacements.get(value, value))
        if not value and rule.get("default") is not None:
            value = str(rule["default"])
        case = rule.get("case")
        require(case in (None, "upper", "lower"), 422, "INVALID_TRANSFORMATION", "大小写转换仅支持 upper 或 lower")
        if case == "upper":
            value = value.upper()
        elif case == "lower":
            value = value.lower()
        output[field] = value
    return output


def _validation_issues(fields, validations):
    require(isinstance(validations, dict) and set(validations) <= {"required", "regex", "allowed", "unique"}, 422, "INVALID_VALIDATION", "校验规则格式不正确")
    required = validations.get("required", [])
    regexes = validations.get("regex", {})
    allowed = validations.get("allowed", {})
    require(isinstance(required, list) and set(required) <= set(FIELDS), 422, "INVALID_VALIDATION", "必填规则包含未知标准字段")
    require(isinstance(regexes, dict) and set(regexes) <= set(FIELDS), 422, "INVALID_VALIDATION", "格式规则包含未知标准字段")
    require(isinstance(allowed, dict) and set(allowed) <= set(FIELDS), 422, "INVALID_VALIDATION", "枚举规则包含未知标准字段")
    issues = []
    for field in required:
        if not str(fields.get(field, "")).strip():
            issues.append({"level": "error", "message": f"必填字段为空：{LABELS.get(field, field)}"})
    for field, pattern in regexes.items():
        require(isinstance(pattern, str) and len(pattern) <= 200, 422, "INVALID_VALIDATION", "正则规则长度不能超过 200")
        try:
            matched = not fields.get(field) or re.fullmatch(pattern, str(fields[field]))
        except re.error as exc:
            raise Problem(422, "INVALID_VALIDATION", "格式校验正则无效") from exc
        if not matched:
            issues.append({"level": "error", "message": f"字段格式不符合配置：{LABELS.get(field, field)}"})
    for field, choices in allowed.items():
        require(isinstance(choices, list) and len(choices) <= 200, 422, "INVALID_VALIDATION", "枚举校验值过多或格式不正确")
        if fields.get(field) and fields[field] not in {str(x) for x in choices}:
            issues.append({"level": "error", "message": f"字段值不在允许范围：{LABELS.get(field, field)}"})
    return issues


def mapped_rows(parsed, sheet, header_row, mapping, catalog=False, transformations=None, validations=None):
    p = preview(parsed, sheet, header_row)
    require(mapping.get("name") in p["headers"], 422, "MAPPING_REQUIRED", "必须指定商品名称列")
    require(all(k in FIELDS and v in p["headers"] for k, v in mapping.items()), 422, "INVALID_MAPPING", "字段映射包含未知列")
    require(len(mapping.values()) == len(set(mapping.values())), 422, "INVALID_MAPPING", "一列不能同时映射多个字段")
    if catalog:
        require(mapping.get("sku") in p["headers"], 422, "MAPPING_REQUIRED", "标准库必须指定内部编号列")
    formula_set = {tuple(x) for x in parsed[sheet]["formulas"]}
    transformations, validations = transformations or {}, validations or {}
    unique_fields = validations.get("unique", [])
    require(isinstance(unique_fields, list) and set(unique_fields) <= set(FIELDS), 422, "INVALID_VALIDATION", "唯一规则包含未知标准字段")
    output, seen, custom_seen = [], set(), set()
    for row_no, row in enumerate(parsed[sheet]["rows"][header_row:], header_row + 1):
        raw = {h: row[i] if i < len(row) else "" for i, h in enumerate(p["headers"])}
        fields = _transform_fields({k: raw[v] for k, v in mapping.items()}, transformations)
        norm = normalize(fields)
        issues = [{"level": "warning", "message": x} for x in norm["issues"]]
        issues.extend(_validation_issues(fields, validations))
        if not norm["name"]:
            issues.append({"level": "error", "message": "商品名称为空"})
        formula_cols = [h for i, h in enumerate(p["headers"]) if (row_no, i) in formula_set]
        if formula_cols:
            issues.append({"level": "error", "message": "含公式单元格，请粘贴为普通值：" + "、".join(formula_cols)})
        sku = fields.get("sku", "").strip()
        if sku and sku in seen:
            issues.append({"level": "error" if catalog else "warning", "message": "商品编号重复"})
        seen.add(sku)
        if unique_fields:
            unique_key = tuple(fields.get(k, "") for k in unique_fields)
            if unique_key in custom_seen:
                issues.append({"level": "error", "message": "配置的唯一字段组合重复"})
            custom_seen.add(unique_key)
        if catalog and not sku:
            issues.append({"level": "error", "message": "内部编号为空"})
        if catalog and norm["issues"]:
            issues.append({"level": "error", "message": "标准库字段存在解析问题或内部冲突"})
        if norm["missing"]:
            issues.append({"level": "warning", "message": "关键字段缺失：" + "、".join(LABELS[k] for k in norm["missing"])})
        output.append({"row_no": row_no, "raw": raw, "normalized": norm, "sku": sku or f"ROW-{row_no:06d}", "generated_sku": not sku, "issues": issues})
    require(len(output) <= (50000 if catalog else 10000), 422, "ROW_LIMIT", "标准库最多 50,000 行，供应商表最多 10,000 行")
    return output
